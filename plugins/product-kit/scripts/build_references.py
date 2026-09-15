#!/usr/bin/env python3
"""Build Word references from the two maintained Markdown template sources.

Requires python-docx. Supports the headings, paragraphs, and pipe tables used by
these sources; this is deliberately not a general Markdown converter.
"""
from pathlib import Path
import argparse
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def build(kind, out=None):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(11.7)
    section.page_height = Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    section.header_distance = section.footer_distance = Inches(.25)
    for element in list(doc.styles.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    doc.styles['Title'].font.size = Pt(28)
    for name, size in [('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(9)
        doc.styles[name].paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = 'PROJECT DOCUMENT TEMPLATES'
    header.style = doc.styles['Normal']
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.text = 'Reusable reference  |  Replace bracketed fields  |  '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)

    lines = (ROOT / 'references' / f'{kind}-template.md').read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            n = len(rows[0])
            proportions = {2: [0.30, .70], 3: [.34, .33, .33],
                           4: [.22, .23, .32, .23], 5: [.10, .26, .14, .25, .25],
                           6: [.13, .28, .18, .14, .10, .17]}[n]
            for col, width in zip(table.columns, proportions):
                col.width = Inches(10.3 * width)
            for row_index, values in enumerate(rows):
                row = table.add_row()
                if row_index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    row._tr.get_or_add_trPr().append(repeat)
                for cell, value, width in zip(row.cells, values, proportions):
                    cell.width = Inches(10.3 * width)
                    cell.text = value
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for side in ['top', 'left', 'bottom', 'right']:
                        border = OxmlElement('w:' + side)
                        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                            border.set(qn('w:' + key), val)
                        borders.append(border)
                    props.append(borders)
                    margins = OxmlElement('w:tcMar')
                    for side in ['top', 'left', 'bottom', 'right']:
                        edge = OxmlElement('w:' + side)
                        edge.set(qn('w:w'), '90')
                        edge.set(qn('w:type'), 'dxa')
                        margins.append(edge)
                    props.append(margins)
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), '233B50' if row_index == 0 else 'F2F5F7' if row_index % 2 else 'FFFFFF')
                    props.append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = row_index == 0
                            run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '000000')
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line[level:].strip()
            doc.add_paragraph(text, 'Title' if level == 1 else f'Heading {min(level - 1, 3)}')
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(re.sub(r'^\d+\. ', '', line), 'List Number')
        else:
            paragraph = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].startswith(('#', '|')) and not re.match(r'^\d+\. ', lines[i + 1]):
                i += 1
                paragraph.append(lines[i].strip())
            doc.add_paragraph(' '.join(paragraph))
        i += 1
    output = (out or ROOT / 'assets') / kind / 'reference.docx'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, help='Optional output root; defaults to skill assets')
    parser.add_argument('--kind', choices=['prd', 'wbs', 'all'], default='all')
    args = parser.parse_args()
    for kind in (['wbs', 'prd'] if args.kind == 'all' else [args.kind]):
        build(kind, args.out)
